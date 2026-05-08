from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(r"D:\claude-code-sourcemap\industrial-ai-report-v9-grounded-examples.pptx")

WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(30, 30, 30)
GRAY = RGBColor(90, 90, 90)
LIGHT = RGBColor(235, 238, 242)
BLUE = RGBColor(232, 242, 255)
BLUE_LINE = RGBColor(70, 130, 220)
ORANGE = RGBColor(255, 242, 226)
ORANGE_LINE = RGBColor(228, 138, 41)
GREEN = RGBColor(232, 247, 236)
GREEN_LINE = RGBColor(68, 160, 87)
PURPLE = RGBColor(240, 235, 251)
PURPLE_LINE = RGBColor(121, 102, 198)
RED = RGBColor(249, 234, 234)
RED_LINE = RGBColor(210, 90, 90)


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, title, bullets, fill=WHITE, line=LIGHT):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = BLACK
    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(12.5)
        p.font.color.rgb = GRAY
        p.space_before = Pt(4)
    return shape


def box(slide, x, y, w, h, text, fill, line, size=15, rounded=True, bold=False):
    t = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(t, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.2)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = BLACK
    return shape


def diamond(slide, x, y, w, h, text):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.DIAMOND, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 249, 233)
    shape.line.color.rgb = ORANGE_LINE
    shape.line.width = Pt(1.2)
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(13)
    r.font.color.rgb = BLACK
    return shape


def line(slide, x1, y1, x2, y2, color=GRAY, dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(1.1)
    c.line.end_arrowhead = True
    if dashed:
        c.line.dash_style = 2
    return c


def title(slide, txt, kicker=None):
    if kicker:
        box(slide, Inches(0.62), Inches(0.18), Inches(1.0), Inches(0.28), kicker, ORANGE, ORANGE_LINE, size=10)
    add_text(slide, Inches(0.65), Inches(0.45), Inches(9.2), Inches(0.45), txt, size=24, bold=True)
    rule = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), Inches(0.98), Inches(9.1), Inches(0.02))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor(220, 220, 220)
    rule.line.fill.background()


def footer(slide, txt):
    add_text(slide, Inches(0.68), Inches(7.0), Inches(8.7), Inches(0.2), txt, size=10, color=GRAY)


def slide_1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    add_text(s, Inches(0.8), Inches(0.9), Inches(5.6), Inches(1.0),
             "工业 AI 故障诊断\n外部项目调研与改造方案", size=26, bold=True)
    add_text(s, Inches(0.82), Inches(2.0), Inches(5.8), Inches(0.7),
             "这次汇报只回答三件事：\n1. 我调研的项目现在到底做到什么程度\n2. 它们最值得学的具体设计点是什么\n3. 这些点如何落到我们的 7 个改进方向", size=15, color=GRAY)
    box(s, Inches(6.35), Inches(1.25), Inches(2.25), Inches(0.65), "调研项目能做到什么程度", ORANGE, ORANGE_LINE, size=15)
    box(s, Inches(6.35), Inches(2.45), Inches(2.25), Inches(0.65), "它们做得好的地方在哪", BLUE, BLUE_LINE, size=15)
    box(s, Inches(6.35), Inches(3.65), Inches(2.25), Inches(0.65), "我们的改进设计是什么", GREEN, GREEN_LINE, size=15)
    line(s, Inches(7.47), Inches(1.9), Inches(7.47), Inches(2.45))
    line(s, Inches(7.47), Inches(3.1), Inches(7.47), Inches(3.65))
    footer(s, "主线不再是泛泛讲 Agent，而是围绕 fault-diagnosis 的具体改造做对标")


def slide_2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title(s, "项目 1：DeerFlow 现在具体能做到什么", "Project")
    add_text(s, Inches(0.72), Inches(1.15), Inches(8.4), Inches(0.42),
             "中心判断：DeerFlow 已经不是“一个 Agent 应用”，而是“一个可编排的 Agent 运行底座”。",
             size=18, bold=True, color=RGBColor(45, 80, 140))
    add_bullets(s, Inches(0.7), Inches(1.7), Inches(4.15), Inches(2.4), "它具体能做到", [
        "统一组织 skills、tools、sub-agents、memory、sandbox",
        "支持 setup wizard、可插拔 skill、长任务协作",
        "支持 tracing、多模型接入、文件系统与隔离执行",
    ], fill=RGBColor(252, 250, 245), line=ORANGE_LINE)
    add_bullets(s, Inches(5.05), Inches(1.7), Inches(4.15), Inches(2.4), "它带来的效果", [
        "复杂任务不再只靠一个大 Prompt 硬扛",
        "新能力可以插进 runtime，而不是持续改主逻辑",
        "长任务的上下文、边界和执行更可控",
    ], fill=RGBColor(245, 249, 255), line=BLUE_LINE)
    add_bullets(s, Inches(0.7), Inches(4.55), Inches(8.5), Inches(1.25), "对 fault-diagnosis 最值得借", [
        "MCP 接入层、Agent workflow 骨架、tool 治理、memory/sub-agent 边界",
        "先做 runtime 分层，再往里面放工业诊断逻辑",
    ], fill=RGBColor(248, 249, 251), line=LIGHT)
    box(s, Inches(0.72), Inches(6.0), Inches(8.45), Inches(0.62),
        "实际例子：如果用户让系统“先查设备告警，再读故障手册，再生成报告”，DeerFlow 这类底座更容易把它拆成可控步骤，而不是全塞进一个大 Prompt。",
        RGBColor(255, 252, 246), ORANGE_LINE, size=12.5)
    footer(s, "一句话：DeerFlow 证明了成熟 Agent 不是大 Prompt + 多工具，而是一个可编排运行时")


def slide_3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title(s, "项目 2：Google ADK 现在具体能做到什么", "Project")
    add_text(s, Inches(0.72), Inches(1.15), Inches(8.4), Inches(0.42),
             "中心判断：Google ADK 的成熟点，不是功能多，而是把“开发、控制、评测、部署”串成一条链。",
             size=18, bold=True, color=RGBColor(45, 80, 140))
    add_bullets(s, Inches(0.7), Inches(1.7), Inches(4.15), Inches(2.4), "它具体能做到", [
        "直接定义 agents、sub_agents、tools、orchestration",
        "接 rich tools、OpenAPI、MCP，不局限于内置能力",
        "支持 tool confirmation、dev UI、eval、Cloud Run / Vertex 部署",
    ], fill=RGBColor(245, 249, 255), line=BLUE_LINE)
    add_bullets(s, Inches(5.05), Inches(1.7), Inches(4.15), Inches(2.4), "它带来的效果", [
        "Agent 不再只是配置，而是工程对象",
        "工具接入、执行确认、评测、部署在同一体系里",
        "多 Agent 层级协作有官方范式，更适合生产化",
    ], fill=RGBColor(252, 250, 245), line=ORANGE_LINE)
    add_bullets(s, Inches(0.7), Inches(4.55), Inches(8.5), Inches(1.25), "对 fault-diagnosis 最值得借", [
        "已有 tool 的标准化治理、MCP 化接入、workflow 的代码化与服务化",
        "如果以后走线上化，这套思路比继续堆脚本更稳",
    ], fill=RGBColor(248, 249, 251), line=LIGHT)
    box(s, Inches(0.72), Inches(6.0), Inches(8.45), Inches(0.62),
        "实际例子：像“生成检修建议前先确认是否允许写工单”这种高风险动作，ADK 的 tool confirmation 思路就比直接放行更适合生产环境。",
        RGBColor(248, 252, 255), BLUE_LINE, size=12.5)
    footer(s, "一句话：Google ADK 证明了成熟 Agent 工程必须把“开发、控制、评测、部署”串成一条链")


def slide_4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title(s, "项目 3：AWS EKS Troubleshooting Agentic AI 现在具体能做到什么", "Project")
    add_text(s, Inches(0.72), Inches(1.05), Inches(8.4), Inches(0.5),
             "中心判断：这是最贴近我们的项目，因为它已经把“垂直排障 Agent”做成了真实工作流。",
             size=18, bold=True, color=RGBColor(45, 80, 140))
    add_bullets(s, Inches(0.72), Inches(1.65), Inches(4.0), Inches(2.45), "它具体能做到", [
        "Slack 里直接发排障请求，系统走 ChatOps 入口",
        "Orchestrator 先分类，再决定是否进入排障 workflow",
        "Memory Agent 查历史案例，K8s Specialist 用 EKS MCP 拉实时状态",
        "最后综合“历史案例 + 当前状态”生成排障建议",
    ], fill=GREEN, line=GREEN_LINE)
    add_bullets(s, Inches(5.0), Inches(1.65), Inches(4.2), Inches(2.45), "它带来的效果", [
        "排障不再只靠静态手册，而能结合实时系统状态",
        "系统能先分类、再编排、再取证，而不是一上来就回答",
        "历史案例能沉淀并反过来帮助后续问题",
    ], fill=RGBColor(247, 250, 255), line=BLUE_LINE)
    # flow
    box(s, Inches(1.0), Inches(4.45), Inches(1.4), Inches(0.55), "Slack / ChatOps", ORANGE, ORANGE_LINE, size=14)
    box(s, Inches(3.0), Inches(4.35), Inches(1.8), Inches(0.72), "Orchestrator\n分类 + 编排", BLUE, BLUE_LINE, size=14)
    box(s, Inches(5.4), Inches(4.35), Inches(1.6), Inches(0.72), "Memory Agent\n历史案例", GREEN, GREEN_LINE, size=13)
    box(s, Inches(5.4), Inches(5.35), Inches(1.6), Inches(0.72), "K8s Specialist\n实时状态", PURPLE, PURPLE_LINE, size=13)
    box(s, Inches(7.6), Inches(4.85), Inches(1.5), Inches(0.55), "排障建议", ORANGE, ORANGE_LINE, size=14)
    line(s, Inches(2.4), Inches(4.72), Inches(3.0), Inches(4.72))
    line(s, Inches(4.8), Inches(4.72), Inches(5.4), Inches(4.72))
    line(s, Inches(4.8), Inches(4.95), Inches(5.4), Inches(5.72))
    line(s, Inches(7.0), Inches(4.72), Inches(7.6), Inches(5.12))
    line(s, Inches(7.0), Inches(5.72), Inches(7.6), Inches(5.12))
    add_bullets(s, Inches(0.7), Inches(5.95), Inches(8.5), Inches(0.8), "对 fault-diagnosis 最值得借", [
        "问题分类、角色协作、历史知识 + 实时状态联合分析、MCP 真正接实时系统",
    ], fill=RGBColor(248, 249, 251), line=LIGHT)
    box(s, Inches(0.72), Inches(6.82), Inches(8.45), Inches(0.42),
        "实际例子：不是只问“某故障码是什么意思”，而是先查相似历史 case，再拉当前实时状态，再综合给出排障建议。", 
        RGBColor(246, 252, 246), GREEN_LINE, size=12)
    footer(s, "一句话：AWS EKS 这个项目最贴近你，因为它本质上就是一个官方排障型 Agent 工作流")


def slide_5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title(s, "项目 4：Giskard 现在具体能做到什么", "Project")
    add_text(s, Inches(0.72), Inches(1.15), Inches(8.4), Inches(0.42),
             "中心判断：Giskard 的价值不在“帮你回答”，而在“帮你证明系统是否真的变好”。",
             size=18, bold=True, color=RGBColor(45, 80, 140))
    add_bullets(s, Inches(0.7), Inches(1.7), Inches(4.15), Inches(2.4), "它具体能做到", [
        "Scenario API 做多轮 agent 测试，不只测单轮输出",
        "Groundedness 检查回答是否真正基于检索证据",
        "LLM-as-judge、回归测试，以及 v2 的 scan / RAGET 能力",
    ], fill=PURPLE, line=PURPLE_LINE)
    add_bullets(s, Inches(5.05), Inches(1.7), Inches(4.15), Inches(2.4), "它带来的效果", [
        "Agent 质量不再靠人工感觉判断",
        "RAG 是否有依据，可以被显式验证",
        "多轮对话与长流程问题可以做回归",
    ], fill=RGBColor(246, 252, 246), line=GREEN_LINE)
    add_bullets(s, Inches(0.7), Inches(4.55), Inches(8.5), Inches(1.25), "对 fault-diagnosis 最值得借", [
        "bad-case、benchmark、groundedness、multi-turn eval",
        "它不会替你做业务逻辑，但能帮你证明 Prompt / RAG / workflow 是否真的更好",
    ], fill=RGBColor(248, 249, 251), line=LIGHT)
    box(s, Inches(0.72), Inches(6.0), Inches(8.45), Inches(0.62),
        "实际例子：同一个故障问法改写成 5 种表达，再检查答案是否都真正引用了检索证据，这就是 Giskard 能帮你做的质量验证。", 
        RGBColor(250, 246, 255), PURPLE_LINE, size=12.5)
    footer(s, "一句话：Giskard 证明了 Agent 做到后面，必须从“能答”走向“能测、能回归、能证明”")


def slide_6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title(s, "工业 AI 现在差在哪里，我们准备怎么专门优化", "Gap")
    headers = ["改进点", "当前差距 / 不足", "专门优化方案"]
    xs = [0.6, 2.35, 6.15]
    ws = [1.65, 3.65, 3.2]
    for x, w, hdr in zip(xs, ws, headers):
        box(s, Inches(x), Inches(1.25), Inches(w), Inches(0.48), hdr, RGBColor(244, 246, 249), LIGHT, size=12.5, rounded=False, bold=True)

    rows = [
        ("MCP", "现状：能力主要还在项目内部函数层；\n差距：还不像 AWS EKS 那样真正接实时系统", "做法：先拆数据 MCP、知识 MCP、动作 MCP；\n优先接设备/告警/工单"),
        ("SSE", "现状：有流式输出，但更像文本流；\n差距：缺少阶段、证据、子 Agent 事件", "做法：补 stage/tool/evidence/error 事件协议；\n前端可视化执行过程"),
        ("Prompt", "现状：Prompt 还偏集中；\n差距：角色边界还不够像垂直排障 workflow", "做法：拆主控/手册/数据/报告四类 Prompt；\n减少一个大 Prompt 包办一切"),
        ("手册 RAG", "现状：更像普通 PDF 检索；\n差距：还没像排障系统那样围绕步骤/风险/故障码组织", "做法：抽故障码、现象、步骤、风险字段；\n给子 Agent 专用知识层"),
        ("FAISS", "现状：能用，但更多是单层向量检索；\n差距：缺少分层路由和验证", "做法：先做手册库/案例库/故障码库分层；\n再用 groundedness 验证效果"),
        ("Workflow", "现状：能跑，但阶段感还不够强；\n差距：不像 AWS EKS 那样先分类、再取证、再汇总", "做法：显式化为分类->取证->分析->汇总->报告"),
        ("Tools", "现状：工具可调，但治理还不统一；\n差距：缺少确认、trace、标准输出契约", "做法：统一 input/output/error/trace；\n高风险动作加 confirmation"),
    ]
    y = 1.73
    fills = [ORANGE, BLUE, GREEN, PURPLE, GREEN, BLUE, ORANGE]
    lines = [ORANGE_LINE, BLUE_LINE, GREEN_LINE, PURPLE_LINE, GREEN_LINE, BLUE_LINE, ORANGE_LINE]
    heights = [0.74, 0.74, 0.74, 0.86, 0.86, 0.86, 0.74]
    for (a, b, c), fill, ln, h in zip(rows, fills, lines, heights):
        box(s, Inches(xs[0]), Inches(y), Inches(ws[0]), Inches(h), a, fill, ln, size=12.5, rounded=False)
        box(s, Inches(xs[1]), Inches(y), Inches(ws[1]), Inches(h), b, RGBColor(250, 250, 250), LIGHT, size=11.2, rounded=False)
        box(s, Inches(xs[2]), Inches(y), Inches(ws[2]), Inches(h), c, RGBColor(247, 250, 255), BLUE_LINE, size=11.2, rounded=False)
        y += h
    footer(s, "这页的重点：不是列方向，而是讲清楚我们和标杆项目的差距，以及对应怎么补")


def slide_7(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title(s, "改造方向 A：把控制面和执行面做扎实", "Design A")
    add_text(s, Inches(0.72), Inches(1.1), Inches(8.6), Inches(0.42),
             "中心判断：方向 A 不是再加功能，而是让系统从“能跑”升级成“可控、可观测、可扩展”。",
             size=18, bold=True, color=RGBColor(45, 80, 140))
    add_bullets(s, Inches(0.7), Inches(1.6), Inches(4.2), Inches(2.2), "这一部分具体讲 4 件事", [
        "MCP：把外部能力从工具函数升级成标准接入层",
        "SSE：把流式输出升级成结构化事件流",
        "Workflow：把隐式推理升级成显式阶段流转",
        "Tool：统一输入、输出、错误、trace 契约",
    ], fill=BLUE, line=BLUE_LINE)
    # flowchart style
    box(s, Inches(5.5), Inches(1.45), Inches(1.6), Inches(0.5), "Context / 身份", ORANGE, ORANGE_LINE, size=13)
    box(s, Inches(7.45), Inches(1.45), Inches(1.8), Inches(0.5), "HTTP /chat/stream", ORANGE, ORANGE_LINE, size=13)
    box(s, Inches(6.0), Inches(2.2), Inches(2.8), Inches(0.72), "中间控制层\nTodo / identity / summary / stage", BLUE, BLUE_LINE, size=13)
    box(s, Inches(6.4), Inches(3.35), Inches(1.2), Inches(0.5), "LLM 推理", ORANGE, ORANGE_LINE, size=13)
    diamond(s, Inches(7.95), Inches(3.2), Inches(0.8), Inches(0.75), "路由")
    box(s, Inches(5.55), Inches(4.35), Inches(1.4), Inches(0.5), "工具执行", BLUE, BLUE_LINE, size=13)
    box(s, Inches(7.95), Inches(4.35), Inches(1.4), Inches(0.5), "SSE 事件", GREEN, GREEN_LINE, size=13)
    box(s, Inches(5.4), Inches(5.15), Inches(4.0), Inches(0.88), "MCP 能力层\n数据 MCP / 知识 MCP / 协作 MCP", PURPLE, PURPLE_LINE, size=13)
    line(s, Inches(6.4), Inches(1.95), Inches(6.9), Inches(2.2), dashed=True)
    line(s, Inches(8.35), Inches(1.95), Inches(8.35), Inches(2.2))
    line(s, Inches(7.4), Inches(2.92), Inches(7.0), Inches(3.35))
    line(s, Inches(7.6), Inches(3.6), Inches(7.95), Inches(3.57))
    line(s, Inches(8.35), Inches(3.95), Inches(8.65), Inches(4.35))
    line(s, Inches(8.0), Inches(3.95), Inches(6.95), Inches(4.35))
    line(s, Inches(6.25), Inches(4.85), Inches(6.25), Inches(5.15))
    box(s, Inches(0.72), Inches(4.25), Inches(4.18), Inches(1.45),
        "具体例子：\n现在：请求会先经过 TodoList / identity_prompt / Summary，再进入模型与 tool loop，按需要调用 sql_inter、query_knowledge_base 等工具后通过 SSE 返回。\n改后：把这条链路进一步显式化为【问题分类】->【数据 MCP】->【知识 MCP】->【事件流回传】->【汇总结论】。",
        RGBColor(248, 249, 251), LIGHT, size=12.2)
    box(s, Inches(0.72), Inches(5.95), Inches(8.55), Inches(0.72),
        "你讲的时候可以这样说：当前已经有 Agent + tool loop，但阶段控制还不够显式；改完后会更像“系统按阶段把事情做完”。",
        RGBColor(255, 252, 246), ORANGE_LINE, size=12.5)
    footer(s, "对应借鉴：DeerFlow 的 runtime 分层、ADK 的工程控制力、AWS EKS 的垂直排障编排")


def slide_8(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title(s, "改造方向 B：把证据面和知识面做扎实", "Design B")
    add_text(s, Inches(0.72), Inches(1.1), Inches(8.6), Inches(0.42),
             "中心判断：方向 B 不是让系统“说得更多”，而是让系统“给出更可信的证据链”。",
             size=18, bold=True, color=RGBColor(45, 80, 140))
    add_bullets(s, Inches(0.7), Inches(1.6), Inches(4.15), Inches(2.3), "这一部分具体讲 3 件事", [
        "不同角色的 Prompt 调优：主控 / 手册 / 数据 / 报告 分开",
        "子 Agent 故障手册 RAG：从 PDF 检索升级成专用知识子系统",
        "FAISS 向量库：先做分层与路由，再考虑换库",
    ], fill=GREEN, line=GREEN_LINE)
    # evidence graph
    box(s, Inches(5.35), Inches(1.55), Inches(1.2), Inches(0.5), "用户问题", ORANGE, ORANGE_LINE, size=13)
    box(s, Inches(6.75), Inches(1.45), Inches(1.6), Inches(0.7), "主控 Agent\n规划与汇总", BLUE, BLUE_LINE, size=13)
    box(s, Inches(5.1), Inches(3.0), Inches(1.5), Inches(0.8), "手册 Agent\n故障码/步骤", GREEN, GREEN_LINE, size=13)
    box(s, Inches(6.95), Inches(3.0), Inches(1.5), Inches(0.8), "数据 Agent\n状态/趋势", PURPLE, PURPLE_LINE, size=13)
    box(s, Inches(8.8), Inches(3.0), Inches(1.2), Inches(0.8), "案例层\n历史经验", ORANGE, ORANGE_LINE, size=12)
    box(s, Inches(6.35), Inches(4.65), Inches(2.5), Inches(0.65), "证据融合：手册 + 数据 + 案例", BLUE, BLUE_LINE, size=13)
    box(s, Inches(6.35), Inches(5.65), Inches(2.5), Inches(0.65), "诊断结论 / 处理建议 / 报告", GREEN, GREEN_LINE, size=13)
    line(s, Inches(6.55), Inches(1.8), Inches(6.75), Inches(1.8))
    line(s, Inches(7.55), Inches(2.15), Inches(5.85), Inches(3.0))
    line(s, Inches(7.55), Inches(2.15), Inches(7.7), Inches(3.0))
    line(s, Inches(7.95), Inches(2.15), Inches(9.25), Inches(3.0))
    line(s, Inches(5.85), Inches(3.8), Inches(6.8), Inches(4.65))
    line(s, Inches(7.7), Inches(3.8), Inches(7.6), Inches(4.65))
    line(s, Inches(9.25), Inches(3.8), Inches(8.2), Inches(4.65))
    line(s, Inches(7.6), Inches(5.3), Inches(7.6), Inches(5.65))
    box(s, Inches(0.72), Inches(4.35), Inches(4.13), Inches(1.3),
        "具体例子：\n现在：系统已经会调用 query_knowledge_base，从 FAISS 检索手册片段，再结合模型组织答案；但手册知识还没有被拆成更清晰的“故障码/步骤/风险”层。\n改后：主控 Agent 先判定场景，手册 Agent 抽取结构化证据，数据 Agent 再判断当前状态是否支持该结论。",
        RGBColor(248, 249, 251), LIGHT, size=12.0)
    box(s, Inches(0.72), Inches(5.92), Inches(8.55), Inches(0.75),
        "你讲的时候可以这样说：当前已经有知识检索，但证据组织还不够细；改完后会更像“先拿结构化证据，再组织结论，再验证结论”。",
        RGBColor(246, 252, 246), GREEN_LINE, size=12.5)
    footer(s, "对应借鉴：AWS EKS 的角色协作与历史知识检索，Giskard 的 groundedness 与 bad-case 验证")


def slide_9(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title(s, "把 7 个改进方向落成一个目标架构", "Target")
    # center architecture
    box(s, Inches(3.85), Inches(1.35), Inches(2.1), Inches(0.72), "Fault Diagnosis Orchestrator", BLUE, BLUE_LINE, size=16)
    box(s, Inches(0.8), Inches(2.6), Inches(2.1), Inches(0.72), "Data MCP\n时序 / 告警 / 资产", ORANGE, ORANGE_LINE, size=14)
    box(s, Inches(3.95), Inches(2.6), Inches(2.1), Inches(0.72), "Knowledge Layer\n手册 / 故障码 / 案例", GREEN, GREEN_LINE, size=14)
    box(s, Inches(7.1), Inches(2.6), Inches(2.1), Inches(0.72), "Action MCP\n工单 / 通知 / 报告", PURPLE, PURPLE_LINE, size=14)
    box(s, Inches(1.3), Inches(4.15), Inches(1.8), Inches(0.72), "Prompt Roles\n主控/手册/数据/报告", BLUE, BLUE_LINE, size=13)
    box(s, Inches(4.05), Inches(4.15), Inches(1.9), Inches(0.72), "RAG Routing\nFAISS 分层路由", GREEN, GREEN_LINE, size=13)
    box(s, Inches(6.95), Inches(4.15), Inches(2.2), Inches(0.72), "SSE Event Bus\n阶段 / 工具 / 证据 / 异常", ORANGE, ORANGE_LINE, size=13)
    box(s, Inches(3.55), Inches(5.5), Inches(2.7), Inches(0.72), "Eval Loop\nbad-case / benchmark / groundedness", RED, RED_LINE, size=13)
    line(s, Inches(4.9), Inches(2.07), Inches(1.85), Inches(2.6))
    line(s, Inches(4.9), Inches(2.07), Inches(5.0), Inches(2.6))
    line(s, Inches(4.9), Inches(2.07), Inches(8.15), Inches(2.6))
    line(s, Inches(1.85), Inches(3.32), Inches(2.2), Inches(4.15))
    line(s, Inches(5.0), Inches(3.32), Inches(5.0), Inches(4.15))
    line(s, Inches(8.15), Inches(3.32), Inches(8.05), Inches(4.15))
    line(s, Inches(5.0), Inches(4.87), Inches(4.9), Inches(5.5))
    add_text(s, Inches(0.8), Inches(6.45), Inches(8.8), Inches(0.42),
             "核心变化：从“一个大 Prompt + 一堆工具”变成“主控编排 + 分层知识 + 标准接入 + 事件观测 + 评测闭环”。",
             size=16, bold=True, color=RGBColor(45, 80, 140))
    footer(s, "这一页是你最后讲“改进后版本长什么样”的关键页")


def slide_10(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title(s, "建议的汇报收束方式", "Close")
    add_bullets(s, Inches(0.8), Inches(1.4), Inches(8.6), Inches(4.0), "你最后可以这样总结", [
        "第一，我不是泛泛看了几个 Agent 项目，而是围绕工业故障诊断最相关的 4 类能力去找标杆。",
        "第二，这些项目已经把 Agent 做到更工程化的阶段：有底座、有垂直排障 workflow、有质量验证。",
        "第三，我们的改造不应该再宽泛发散，而应该明确围绕 7 个方向推进。",
        "第四，这 7 个方向并不是分散的功能点，而是一套完整的目标架构。",
    ], fill=RGBColor(248, 249, 251), line=LIGHT)
    box(s, Inches(1.2), Inches(5.75), Inches(2.0), Inches(0.65), "调研项目做到什么程度", ORANGE, ORANGE_LINE, size=15)
    box(s, Inches(3.9), Inches(5.75), Inches(2.0), Inches(0.65), "我们借哪些具体设计", BLUE, BLUE_LINE, size=15)
    box(s, Inches(6.6), Inches(5.75), Inches(2.0), Inches(0.65), "改进后的系统长什么样", GREEN, GREEN_LINE, size=15)
    line(s, Inches(3.2), Inches(6.08), Inches(3.9), Inches(6.08))
    line(s, Inches(5.9), Inches(6.08), Inches(6.6), Inches(6.08))
    footer(s, "这页不是技术页，而是帮你把整场汇报收成一个清楚的闭环")


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    for fn in [slide_1, slide_2, slide_3, slide_4, slide_5, slide_6, slide_7, slide_8, slide_9, slide_10]:
        fn(prs)
    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
