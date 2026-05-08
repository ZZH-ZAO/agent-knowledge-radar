from pathlib import Path

from pptx import Presentation


PPT_PATH = Path(r"D:\claude-code-sourcemap\industrial-ai-report-v9-grounded-examples.pptx")
OUT_PATH = Path(r"D:\claude-code-sourcemap\industrial-ai-report-v9-grounded-examples-harness.pptx")


def replace_text(shape, old, new):
    if not hasattr(shape, "text_frame"):
        return False
    text = shape.text
    if old not in text:
        return False
    text = text.replace(old, new)
    tf = shape.text_frame
    paragraphs = list(tf.paragraphs)
    first_run = None
    first_para = None
    for p in paragraphs:
        if p.runs:
            first_run = p.runs[0]
            first_para = p
            break
    while len(tf.paragraphs) > 1:
        p = tf.paragraphs[-1]
        p._element.getparent().remove(p._element)
    tf.paragraphs[0].clear()
    run = tf.paragraphs[0].add_run()
    run.text = text
    if first_run is not None:
        run.font.name = first_run.font.name
        run.font.size = first_run.font.size
        run.font.bold = first_run.font.bold
        if first_run.font.color is not None and first_run.font.color.type is not None:
            try:
                run.font.color.rgb = first_run.font.color.rgb
            except Exception:
                pass
        tf.paragraphs[0].alignment = first_para.alignment
    return True


def main():
    prs = Presentation(str(PPT_PATH))
    slide = prs.slides[1]  # DeerFlow page

    replacements = [
        (
            "中心判断：DeerFlow 已经不是“一个 Agent 应用”，而是“一个可编排的 Agent 运行底座”。",
            "中心判断：DeerFlow 已经不是普通 Agent 应用，而是一个 super agent harness，也就是可复用的 Agent 执行底座。"
        ),
        (
            "它具体能做到 / 统一组织 skills、tools、sub-agents、memory、sandbox / 支持 setup wizard、可插拔 skill、长任务协作 / 支持 tracing、多模型接入、文件系统与隔离执行",
            "它具体能做到 / 统一组织 skills、tools、sub-agents、memory、sandbox / 这些能力运行在同一套 runtime 里 / harness 负责把这套 runtime 封装成可复用底座"
        ),
        (
            "对 fault-diagnosis 最值得借 / MCP 接入层、Agent workflow 骨架、tool 治理、memory/sub-agent 边界 / 先做 runtime 分层，再往里面放工业诊断逻辑",
            "对 fault-diagnosis 最值得借 / 不是单个功能，而是 harness 思路 / 先把 runtime 分层和 Agent 底座搭起来，再往里面放工业诊断逻辑"
        ),
        (
            "一句话：DeerFlow 证明了成熟 Agent 不是大 Prompt + 多工具，而是一个可编排运行时",
            "一句话：DeerFlow 证明了成熟 Agent 不是大 Prompt + 多工具，而是一个可复用的 Agent harness。"
        ),
    ]

    for shape in slide.shapes:
        for old, new in replacements:
            replace_text(shape, old, new)

    prs.save(str(OUT_PATH))
    print(OUT_PATH)


if __name__ == "__main__":
    main()
