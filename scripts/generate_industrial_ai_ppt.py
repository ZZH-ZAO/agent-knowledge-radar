from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(r"D:\claude-code-sourcemap\industrial-ai-report-v3.pptx")


BLUE = RGBColor(217, 232, 252)
BLUE_LINE = RGBColor(62, 120, 216)
ORANGE = RGBColor(254, 235, 209)
ORANGE_LINE = RGBColor(230, 137, 41)
GREEN = RGBColor(224, 245, 227)
GREEN_LINE = RGBColor(65, 160, 84)
PURPLE = RGBColor(232, 224, 248)
PURPLE_LINE = RGBColor(118, 97, 196)
GRAY = RGBColor(80, 80, 80)
LIGHT_GRAY = RGBColor(110, 110, 110)
RED = RGBColor(244, 220, 220)
RED_LINE = RGBColor(210, 90, 90)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(35, 35, 35)


def set_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def textbox(slide, x, y, w, h, text, size=18, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def bullet_box(slide, x, y, w, h, title, bullets):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = BLACK

    for item in bullets:
        p = tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(14)
        p.font.color.rgb = GRAY
        p.space_before = Pt(5)
    return box


def add_round_box(slide, x, y, w, h, text, fill_rgb, line_rgb, size=16, bold=False):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    shape.line.color.rgb = line_rgb
    shape.line.width = Pt(1.25)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK
    return shape


def add_rect(slide, x, y, w, h, text, fill_rgb, line_rgb, size=16, bold=False):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    shape.line.color.rgb = line_rgb
    shape.line.width = Pt(1.2)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK
    return shape


def add_diamond(slide, x, y, w, h, text):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.DIAMOND, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 248, 230)
    shape.line.color.rgb = ORANGE_LINE
    shape.line.width = Pt(1.2)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(15)
    run.font.color.rgb = BLACK
    return shape


def add_line(slide, x1, y1, x2, y2, color=GRAY, dashed=False, arrow=True):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(1.1)
    if dashed:
        line.line.dash_style = 2
    if arrow:
        line.line.end_arrowhead = True
    return line


def title_and_rule(slide, title, subtitle=None):
    textbox(slide, Inches(0.6), Inches(0.25), Inches(10.5), Inches(0.55), title, size=24, bold=True)
    if subtitle:
        textbox(slide, Inches(0.65), Inches(0.78), Inches(10.1), Inches(0.32), subtitle, size=10, color=LIGHT_GRAY)
    rule = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.6), Inches(1.08), Inches(10.2), Inches(0.02))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor(220, 220, 220)
    rule.line.fill.background()


def add_card(slide, x, y, w, h, title, lines, fill_rgb=WHITE, line_rgb=RGBColor(220, 220, 220)):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    shape.line.color.rgb = line_rgb
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = BLACK
    for line in lines:
        p = tf.add_paragraph()
        p.text = line
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(12.5)
        p.font.color.rgb = GRAY
        p.space_before = Pt(4)
    return shape


def add_footer(slide, text):
    textbox(slide, Inches(0.65), Inches(7.0), Inches(10.0), Inches(0.25), text, size=10, color=LIGHT_GRAY)


def section_pill(slide, text, x=0.65, y=0.18, w=1.15, color_fill=RGBColor(255, 245, 230), color_line=ORANGE_LINE):
    return add_round_box(slide, Inches(x), Inches(y), Inches(w), Inches(0.32), text, color_fill, color_line, size=10)


def slide1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    textbox(slide, Inches(0.85), Inches(0.75), Inches(6.2), Inches(1.1),
            "工业 AI 故障诊断系统\n外部项目调研与升级方向思考", size=24, bold=True)
    textbox(slide, Inches(0.9), Inches(2.0), Inches(5.8), Inches(0.6),
            "围绕 MCP、SSE、Prompt、RAG、FAISS、Workflow、Tool 治理的系统化改造",
            size=12, color=LIGHT_GRAY)
    add_round_box(slide, Inches(6.4), Inches(1.3), Inches(1.7), Inches(0.65), "外部项目调研", ORANGE, ORANGE_LINE, size=15)
    add_round_box(slide, Inches(6.4), Inches(2.45), Inches(1.7), Inches(0.65), "能力抽象对标", BLUE, BLUE_LINE, size=15)
    add_round_box(slide, Inches(6.4), Inches(3.6), Inches(1.7), Inches(0.65), "工业 AI 升级", GREEN, GREEN_LINE, size=15)
    add_line(slide, Inches(7.25), Inches(1.95), Inches(7.25), Inches(2.45), color=GRAY)
    add_line(slide, Inches(7.25), Inches(3.1), Inches(7.25), Inches(3.6), color=GRAY)
    textbox(slide, Inches(0.9), Inches(5.25), Inches(8.2), Inches(0.8),
            "调研目标不是“找几个 demo”，而是回答两个问题：\n1. 优秀 Agent 项目已经把系统做到什么程度\n2. 我们的工业 AI 下一步最值得补哪两层能力",
            size=16, color=GRAY)
    add_footer(slide, "主线：先看标杆，再抽能力，最后落到 fault-diagnosis 的升级路线")


def slide2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title_and_rule(slide, "为什么要看外部优秀项目")
    bullet_box(slide, Inches(0.8), Inches(1.55), Inches(4.2), Inches(3.2), "当前工业 AI 常见问题", [
        "系统能跑通，但能力都堆在主工程里",
        "有诊断结果，但证据链和验证链偏弱",
        "有 tool 和 RAG，但缺少统一治理和观测",
    ])
    add_round_box(slide, Inches(5.6), Inches(2.0), Inches(4.2), Inches(2.05),
                  "不是照搬项目，而是回答两个问题\n\n1. 成熟 Agent 系统已经把哪些层做厚了\n2. 我们工业 AI 下一步最值得补哪两层能力",
                  RGBColor(248, 249, 251), RGBColor(210, 210, 210), size=16)
    textbox(slide, Inches(0.8), Inches(5.35), Inches(8.9), Inches(0.45),
            "这次调研不是“找灵感”，而是“找标杆”。", size=20, bold=True, color=RGBColor(45, 80, 140))
    add_footer(slide, "判断重点：我们缺的是单点功能，还是系统层能力")


def slide3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    section_pill(slide, "Survey", x=0.7, y=0.18, w=0.95)
    title_and_rule(slide, "本次重点参考的外部项目")
    card_w, card_h = Inches(4.5), Inches(2.0)
    add_card(slide, Inches(0.7), Inches(1.4), card_w, card_h, "deer-flow", [
        "类型：Agent 底座 / 运行时系统",
        "来源：bytedance/deer-flow",
        "关注：skills、memory、sub-agent、runtime 治理",
        "启发：能力接入层与执行编排层要分开设计",
    ], fill_rgb=RGBColor(252, 250, 245), line_rgb=RGBColor(232, 204, 160))
    add_card(slide, Inches(5.5), Inches(1.4), card_w, card_h, "Google adk-python", [
        "类型：代码优先 Agent 框架",
        "来源：google/adk-python",
        "关注：tools、MCP、多 Agent、eval、deploy",
        "启发：开发、评测、部署要一体化思考",
    ], fill_rgb=RGBColor(245, 249, 255), line_rgb=RGBColor(182, 205, 238))
    add_card(slide, Inches(0.7), Inches(3.8), card_w, card_h, "giskard", [
        "类型：LLM / RAG / Agent 评测框架",
        "来源：Giskard-AI/giskard",
        "关注：multi-turn eval、RAG eval、judge",
        "启发：工业 AI 需要系统化质量评测",
    ], fill_rgb=RGBColor(246, 252, 246), line_rgb=RGBColor(179, 219, 185))
    add_card(slide, Inches(5.5), Inches(3.8), card_w, card_h, "repomind", [
        "类型：验证导向的 Agent 产品案例",
        "来源：社区开源项目",
        "关注：过程透明、状态记录、verification-first",
        "启发：结论之外，还要让用户看到过程和依据",
    ], fill_rgb=RGBColor(250, 246, 255), line_rgb=RGBColor(207, 188, 233))
    add_footer(slide, "底座型项目回答“怎么做得更可扩展”，质量型项目回答“怎么做得更可信”")


def slide4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    section_pill(slide, "Runtime", x=0.7, y=0.18, w=1.05, color_fill=RGBColor(245, 249, 255), color_line=BLUE_LINE)
    title_and_rule(slide, "参考项目展开（一）：DeerFlow 与 ADK 为什么值得重点借鉴")
    add_round_box(slide, Inches(0.9), Inches(1.35), Inches(8.5), Inches(0.7),
                  "真正值得学的不是某个功能点，而是它们怎么把 Agent 做成“可演进的底座”。",
                  RGBColor(248, 249, 251), RGBColor(210, 210, 210), size=19, bold=True)

    add_card(slide, Inches(0.7), Inches(2.35), Inches(4.35), Inches(2.65), "DeerFlow", [
        "它是什么：通用 Super Agent Harness + 参考应用",
        "它解决什么：多工具、多子 Agent、多运行环境如何组织",
        "它为什么更强：把 middleware、memory、skill、sandbox、MCP 放进统一 runtime",
        "工业 AI 借什么：能力接入层、运行时治理、产品层与 runtime 层分离",
        "边界：它偏通用平台，不能替代工业场景自己的诊断 workflow",
    ], fill_rgb=RGBColor(252, 250, 245), line_rgb=RGBColor(232, 204, 160))
    add_card(slide, Inches(5.0), Inches(2.35), Inches(4.35), Inches(2.65), "Google ADK", [
        "它是什么：大公司官方 code-first Agent 框架",
        "它解决什么：tools、MCP、多 Agent、eval、deploy 如何一体化",
        "它为什么更强：不是只跑起来，而是强调受控执行和后续部署",
        "工业 AI 借什么：tool confirmation / HITL、multi-agent hierarchy、治理思维",
        "边界：它是通用框架，不会自动给你工业证据链",
    ], fill_rgb=RGBColor(245, 249, 255), line_rgb=RGBColor(182, 205, 238))

    add_round_box(slide, Inches(1.1), Inches(5.45), Inches(2.25), Inches(0.62), "Runtime-first", ORANGE, ORANGE_LINE, size=15)
    add_round_box(slide, Inches(3.85), Inches(5.45), Inches(2.25), Inches(0.62), "Governance-first", BLUE, BLUE_LINE, size=15)
    add_round_box(slide, Inches(6.6), Inches(5.45), Inches(2.25), Inches(0.62), "Platform-thinking", GREEN, GREEN_LINE, size=15)
    textbox(slide, Inches(0.9), Inches(6.3), Inches(8.7), Inches(0.45),
            "对工业 AI 的直接启发：不要继续只补功能，而要开始补控制面、接入面和运行时骨架。",
            size=16, color=RGBColor(45, 80, 140), bold=True)
    add_footer(slide, "这页你可以重点讲：为什么我们现在最该补的是控制面与接入面")


def slide5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    section_pill(slide, "Quality", x=0.7, y=0.18, w=0.95, color_fill=RGBColor(246, 252, 246), color_line=GREEN_LINE)
    title_and_rule(slide, "参考项目展开（二）：Giskard 与 RepoMind 为什么值得重点借鉴")
    add_round_box(slide, Inches(0.9), Inches(1.35), Inches(8.5), Inches(0.7),
                  "真正值得学的不是“怎么答得更像”，而是“怎么证明答得更对”。",
                  RGBColor(248, 249, 251), RGBColor(210, 210, 210), size=19, bold=True)
    add_card(slide, Inches(0.7), Inches(2.35), Inches(4.35), Inches(2.65), "Giskard", [
        "它是什么：面向 Agent / RAG 的评测与脆弱性扫描框架",
        "它解决什么：把一次性感觉判断升级成多轮、可复跑、可回归评测",
        "它为什么更强：scenario API、LLM-as-judge、groundedness、RAG evaluation",
        "工业 AI 借什么：bad-case 沉淀、回归评测、多轮稳定性验证",
        "边界：它偏质量系统，不会替你写业务 workflow",
    ], fill_rgb=RGBColor(246, 252, 246), line_rgb=RGBColor(179, 219, 185))
    add_card(slide, Inches(5.0), Inches(2.35), Inches(4.35), Inches(2.65), "RepoMind", [
        "它是什么：产品化的仓库理解与安全验证 Agent 平台",
        "它解决什么：怎么把“能答”做成“稳定地答、带证据地答”",
        "它为什么更强：结构化事件、状态记录、缓存、验证门、结果分享",
        "工业 AI 借什么：verification-first、过程透明、证据优先",
        "边界：它是代码仓库场景，但验证闭环思想很适合工业 AI",
    ], fill_rgb=RGBColor(250, 246, 255), line_rgb=RGBColor(207, 188, 233))

    add_round_box(slide, Inches(1.1), Inches(5.45), Inches(2.25), Inches(0.62), "Evidence-first", GREEN, GREEN_LINE, size=15)
    add_round_box(slide, Inches(3.85), Inches(5.45), Inches(2.25), Inches(0.62), "Verification-first", PURPLE, PURPLE_LINE, size=15)
    add_round_box(slide, Inches(6.6), Inches(5.45), Inches(2.25), Inches(0.62), "Trust-building", ORANGE, ORANGE_LINE, size=15)
    textbox(slide, Inches(0.9), Inches(6.3), Inches(8.7), Inches(0.45),
            "对工业 AI 的直接启发：不要只给结论，而要给证据链、验证门和持续回归机制。",
            size=16, color=RGBColor(45, 80, 140), bold=True)
    add_footer(slide, "这页你可以重点讲：为什么我们现在最该补的是证据面与质量面")


def slide6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title_and_rule(slide, "这些项目已经把 Agent 做到什么程度")
    x0 = Inches(0.7)
    y0 = Inches(1.45)
    widths = [Inches(1.8), Inches(3.3), Inches(4.8)]
    headers = ["项目", "已做厚的层", "对工业 AI 的启发"]
    x = x0
    for idx, hdr in enumerate(headers):
        add_rect(slide, x, y0, widths[idx], Inches(0.5), hdr, RGBColor(242, 245, 249), RGBColor(205, 210, 220), size=14, bold=True)
        x += widths[idx]
    rows = [
        ("deer-flow", "skills、memory、sub-agent、runtime 治理", "工业 AI 不能继续把所有能力都堆在主工程里"),
        ("adk-python", "tools、MCP、多 Agent、eval、deploy", "工业 AI 要把接入、执行、评测、部署串起来考虑"),
        ("giskard", "multi-turn testing、RAG eval、judge", "工业 AI 要从“能答”走向“能评”"),
        ("repomind", "过程透明、状态管理、验证闭环", "工业 AI 要给证据和过程，而不只给结论"),
    ]
    y = Inches(1.95)
    fills = [RGBColor(255, 252, 246), RGBColor(247, 250, 255), RGBColor(247, 252, 247), RGBColor(251, 248, 255)]
    for i, row in enumerate(rows):
        x = x0
        for j, cell in enumerate(row):
            add_rect(slide, x, y, widths[j], Inches(0.68), cell, fills[i], RGBColor(225, 225, 225), size=12)
            x += widths[j]
        y += Inches(0.68)
    add_round_box(slide, Inches(1.25), Inches(5.1), Inches(3.0), Inches(0.8), "Control / Runtime Layer\n接入、编排、治理、观测", BLUE, BLUE_LINE, size=16)
    add_round_box(slide, Inches(6.1), Inches(5.1), Inches(3.0), Inches(0.8), "Evidence / Quality Layer\n证据、评测、回归、验证", GREEN, GREEN_LINE, size=16)
    add_round_box(slide, Inches(3.8), Inches(6.15), Inches(2.8), Inches(0.55), "Industrial AI Next Priorities", ORANGE, ORANGE_LINE, size=14, bold=True)
    add_line(slide, Inches(4.25), Inches(5.9), Inches(4.95), Inches(6.15))
    add_line(slide, Inches(7.6), Inches(5.9), Inches(5.45), Inches(6.15))
    add_footer(slide, "优秀项目的重点已经不是“多调几个工具”，而是把 Agent 做成分层、可控、可验证的系统")


def slide7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title_and_rule(slide, "方向一：补工业 AI 的控制面与接入面", "主要借鉴：deer-flow + Google adk-python")
    # Flowchart like sample
    add_round_box(slide, Inches(1.2), Inches(1.2), Inches(1.9), Inches(0.55), "Context / 身份", RGBColor(255, 247, 225), ORANGE_LINE, size=14)
    add_round_box(slide, Inches(6.0), Inches(1.2), Inches(2.2), Inches(0.55), "HTTP GET /chat/stream", RGBColor(255, 244, 224), ORANGE_LINE, size=14)
    mid = add_rect(slide, Inches(2.1), Inches(2.0), Inches(4.8), Inches(0.75),
                   "控制中间层\nTodoList / identity_prompt / Summary / Stage Router",
                   BLUE, BLUE_LINE, size=15, bold=False)
    add_rect(slide, Inches(3.0), Inches(3.35), Inches(1.55), Inches(0.55), "LLM 推理", ORANGE, ORANGE_LINE, size=15)
    add_rect(slide, Inches(2.8), Inches(5.0), Inches(2.0), Inches(0.7), "工具执行", BLUE, BLUE_LINE, size=15)
    add_round_box(slide, Inches(5.95), Inches(5.0), Inches(1.95), Inches(0.55), "SSE 事件流", GREEN, GREEN_LINE, size=15)
    add_rect(slide, Inches(2.2), Inches(5.95), Inches(3.2), Inches(1.1),
             "MCP 能力层\nData MCP / Knowledge MCP / Action MCP",
             RGBColor(236, 243, 255), BLUE_LINE, size=15)
    add_rect(slide, Inches(2.2), Inches(7.0), Inches(3.2), Inches(0.35),
             "", PURPLE, PURPLE_LINE, size=12)
    textbox(slide, Inches(2.35), Inches(7.02), Inches(2.9), Inches(0.25),
            "外部系统：时序 / 告警 / 手册 / 工单 / 报告", size=11, color=BLACK, align=PP_ALIGN.CENTER)
    add_diamond(slide, Inches(5.15), Inches(3.45), Inches(1.0), Inches(0.8), "路由")
    add_rect(slide, Inches(4.1), Inches(4.45), Inches(1.5), Inches(0.55), "LLM API", PURPLE, PURPLE_LINE, size=15)
    add_round_box(slide, Inches(6.1), Inches(2.95), Inches(1.35), Inches(0.65), "状态持久化", GREEN, GREEN_LINE, size=13)
    add_rect(slide, Inches(6.18), Inches(3.8), Inches(1.2), Inches(0.75), "PostgreSQL", RGBColor(240, 252, 240), GREEN_LINE, size=13)

    add_line(slide, Inches(2.15), Inches(1.75), Inches(2.5), Inches(2.0), color=GRAY, dashed=True)
    add_line(slide, Inches(7.1), Inches(1.75), Inches(7.1), Inches(2.0), color=GRAY)
    add_line(slide, Inches(4.5), Inches(2.75), Inches(3.8), Inches(3.35))
    add_line(slide, Inches(4.75), Inches(2.75), Inches(5.6), Inches(3.45))
    add_line(slide, Inches(3.8), Inches(3.9), Inches(3.8), Inches(5.0))
    add_line(slide, Inches(5.95), Inches(4.05), Inches(6.95), Inches(5.0))
    add_line(slide, Inches(3.8), Inches(5.7), Inches(3.8), Inches(5.95))
    add_line(slide, Inches(6.18), Inches(3.6), Inches(6.78), Inches(3.8), dashed=True)
    add_line(slide, Inches(4.85), Inches(3.62), Inches(4.85), Inches(4.45), dashed=True)
    textbox(slide, Inches(4.65), Inches(4.12), Inches(0.9), Inches(0.25), "调用模型", size=10, color=GREEN_LINE, align=PP_ALIGN.CENTER)
    textbox(slide, Inches(2.1), Inches(4.5), Inches(0.9), Inches(0.25), "ToolMessage", size=10, color=GREEN_LINE)
    textbox(slide, Inches(0.7), Inches(5.7), Inches(1.15), Inches(0.9), "改造重点：\nMCP\nSSE\nWorkflow\nTool 契约", size=13, color=BLACK)
    add_footer(slide, "目标：让工业 AI 从“能跑通”升级成“可扩展、可观测、可治理的系统”")


def slide8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title_and_rule(slide, "方向一展开：MCP、SSE、Workflow、Tool 为什么值得优先做")
    titles = ["MCP", "SSE", "Workflow", "Tool 治理"]
    fills = [ORANGE, BLUE, GREEN, PURPLE]
    lines = [ORANGE_LINE, BLUE_LINE, GREEN_LINE, PURPLE_LINE]
    contents = [
        ["目标：降低异构系统接入成本", "更好：主控 Agent 面向统一能力接口", "更需要：开始接第二、第三个系统时", "欠缺：小系统过早全面 MCP 化有成本"],
        ["目标：让执行过程可观测", "更好：前端可看到阶段、工具、证据、子 Agent 状态", "更需要：复杂诊断、多步执行场景", "欠缺：事件定义混乱会增加复杂度"],
        ["目标：让复杂诊断有稳定骨架", "更好：避免模型直接跳到不可靠结论", "更需要：故障排查、报告生成、多人协作", "欠缺：过度流程化会压制灵活性"],
        ["目标：让工具层可复用、可维护、可审计", "更好：工具越多越需要统一契约", "更需要：已有十几个 tool 且还会扩展时", "欠缺：前期设计成本会上升"],
    ]
    positions = [(0.75, 1.45), (5.55, 1.45), (0.75, 4.2), (5.55, 4.2)]
    for idx, (x, y) in enumerate(positions):
        add_card(slide, Inches(x), Inches(y), Inches(4.05), Inches(2.15), titles[idx], contents[idx], fill_rgb=fills[idx], line_rgb=lines[idx])
    add_footer(slide, "方向一解决的是“系统能不能稳稳长大”，而不只是“这次能不能跑通”")


def slide9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title_and_rule(slide, "方向二：补工业 AI 的证据面与质量面", "主要借鉴：giskard + repomind")
    add_round_box(slide, Inches(0.8), Inches(1.55), Inches(1.3), Inches(0.55), "用户问题", ORANGE, ORANGE_LINE, size=14)
    add_rect(slide, Inches(2.45), Inches(1.45), Inches(2.0), Inches(0.78), "主控诊断 Agent", BLUE, BLUE_LINE, size=16)
    add_rect(slide, Inches(6.0), Inches(1.45), Inches(2.0), Inches(0.78), "证据汇总 / 冲突处理", GREEN, GREEN_LINE, size=15)
    add_round_box(slide, Inches(8.6), Inches(1.55), Inches(1.35), Inches(0.55), "结论 / 报告", PURPLE, PURPLE_LINE, size=14)
    add_card(slide, Inches(1.55), Inches(3.1), Inches(2.1), Inches(1.3), "Data Agent", [
        "读取时序、状态、异常趋势",
        "更偏事实分析，不直接下最终结论",
    ], fill_rgb=RGBColor(246, 250, 255), line_rgb=BLUE_LINE)
    add_card(slide, Inches(4.1), Inches(3.1), Inches(2.1), Inches(1.3), "Manual Agent", [
        "读取故障手册、故障码、处理步骤",
        "强调引用依据，不额外猜测",
    ], fill_rgb=RGBColor(255, 250, 242), line_rgb=ORANGE_LINE)
    add_card(slide, Inches(6.65), Inches(3.1), Inches(2.1), Inches(1.3), "Case Agent", [
        "对照历史案例和相似问题",
        "提供经验型证据",
    ], fill_rgb=RGBColor(246, 252, 246), line_rgb=GREEN_LINE)
    add_line(slide, Inches(2.1), Inches(1.83), Inches(2.45), Inches(1.83))
    add_line(slide, Inches(4.45), Inches(1.83), Inches(6.0), Inches(1.83))
    add_line(slide, Inches(8.0), Inches(1.83), Inches(8.6), Inches(1.83))
    add_line(slide, Inches(3.45), Inches(2.23), Inches(2.6), Inches(3.1))
    add_line(slide, Inches(3.45), Inches(2.23), Inches(5.15), Inches(3.1))
    add_line(slide, Inches(7.0), Inches(3.1), Inches(7.0), Inches(2.23))
    add_line(slide, Inches(7.0), Inches(2.23), Inches(7.0), Inches(1.45))

    add_rect(slide, Inches(1.0), Inches(5.35), Inches(1.55), Inches(0.55), "当前输出", RGBColor(248, 248, 248), RGBColor(200, 200, 200), size=14)
    add_rect(slide, Inches(2.75), Inches(5.35), Inches(1.55), Inches(0.55), "指标评测", RGBColor(248, 248, 248), RGBColor(200, 200, 200), size=14)
    add_rect(slide, Inches(4.5), Inches(5.35), Inches(1.55), Inches(0.55), "bad-case", RED, RED_LINE, size=14)
    add_rect(slide, Inches(6.25), Inches(5.35), Inches(1.55), Inches(0.55), "benchmark", RGBColor(248, 248, 248), RGBColor(200, 200, 200), size=14)
    add_rect(slide, Inches(8.0), Inches(5.35), Inches(1.65), Inches(0.55), "Prompt / RAG / Tool 优化", RGBColor(248, 248, 248), RGBColor(200, 200, 200), size=13)
    for start in [2.55, 4.3, 6.05, 7.8]:
        add_line(slide, Inches(start), Inches(5.62), Inches(start + 0.2), Inches(5.62))
    add_footer(slide, "方向二解决的不是“会不会回答”，而是“这个结论有没有证据、能不能持续验证、出了错能不能反哺优化”")


def slide10(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title_and_rule(slide, "方向二展开：Prompt、手册 RAG、FAISS、质量评估怎么升级")
    titles = ["角色化 Prompt", "子 Agent 故障手册 RAG", "FAISS", "质量评估"]
    fills = [BLUE, ORANGE, GREEN, PURPLE]
    lines = [BLUE_LINE, ORANGE_LINE, GREEN_LINE, PURPLE_LINE]
    contents = [
        ["目标：把职责边界写清楚", "升级：主控 / 手册 / 数据 / 报告 Agent 分开", "更好：更容易测，也更容易定位 bad-case", "欠缺：角色切太细会增加协调成本"],
        ["目标：把 PDF 检索升级成专用知识子系统", "升级：抽取故障码、现象、原因、步骤、风险提示", "更好：减少“找到段落却抓不到关键步骤”", "欠缺：前期知识加工成本更高"],
        ["目标：先做知识分层和检索路由", "升级：手册库、故障码库、案例库分开再汇总", "更好：比一上来换库更符合当前阶段", "欠缺：规模继续增长后仍可能需要更强底座"],
        ["目标：从“感觉答得可以”走向“可回归、可度量”", "升级：bad-case、benchmark、groundedness、多轮评测", "更好：能证明系统是否真的进步", "欠缺：需要持续维护评测集"],
    ]
    positions = [(0.75, 1.45), (5.55, 1.45), (0.75, 4.2), (5.55, 4.2)]
    for idx, (x, y) in enumerate(positions):
        add_card(slide, Inches(x), Inches(y), Inches(4.05), Inches(2.2), titles[idx], contents[idx], fill_rgb=fills[idx], line_rgb=lines[idx])
    add_footer(slide, "方向二解决的是“系统能不能给出可信结论”，而不只是“系统能不能给出结论”")


def slide11(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title_and_rule(slide, "工业 AI 下一步路线图")
    add_card(slide, Inches(1.1), Inches(2.0), Inches(3.6), Inches(2.6), "Phase P0：控制面与接入面", [
        "MCP 分类接入",
        "SSE 事件协议",
        "Workflow 状态化",
        "Tool 契约化",
    ], fill_rgb=RGBColor(245, 249, 255), line_rgb=BLUE_LINE)
    add_card(slide, Inches(6.1), Inches(2.0), Inches(3.6), Inches(2.6), "Phase P1：证据面与质量面", [
        "角色 Prompt 拆分",
        "手册 RAG 分层",
        "FAISS 路由化",
        "bad-case / benchmark / eval",
    ], fill_rgb=RGBColor(246, 252, 246), line_rgb=GREEN_LINE)
    add_line(slide, Inches(4.8), Inches(3.3), Inches(6.1), Inches(3.3), color=ORANGE_LINE)
    add_round_box(slide, Inches(4.9), Inches(3.0), Inches(1.1), Inches(0.55), "先补底座\n再补质量", ORANGE, ORANGE_LINE, size=13)
    textbox(slide, Inches(1.2), Inches(5.45), Inches(8.7), Inches(0.55),
            "实施建议：先做 P0，再做 P1。控制面和接入面是地基；证据面和质量面是系统从“能用”走向“可信”的关键。",
            size=15, color=GRAY)
    add_footer(slide, "先补底座，再补质量，不求一步做全")


def slide12(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title_and_rule(slide, "总结")
    textbox(slide, Inches(1.0), Inches(1.8), Inches(8.7), Inches(0.65),
            "优秀 Agent 项目的重点，已经不是让模型更会说，\n而是让系统更分层、更可控、更可验证。", size=22, bold=True, color=RGBColor(40, 70, 120))
    bullet_box(slide, Inches(1.1), Inches(3.1), Inches(5.3), Inches(1.7), "工业 AI 下一步最值得补的两层能力", [
        "控制面与接入面",
        "证据面与质量面",
    ])
    add_round_box(slide, Inches(0.95), Inches(5.5), Inches(2.25), Inches(0.7), "Runnable Demo\n能跑的原型", ORANGE, ORANGE_LINE, size=16)
    add_round_box(slide, Inches(4.0), Inches(5.5), Inches(2.25), Inches(0.7), "Managed System\n可治理的系统", BLUE, BLUE_LINE, size=16)
    add_round_box(slide, Inches(7.05), Inches(5.5), Inches(2.25), Inches(0.7), "Trusted Platform\n可持续演进的平台", GREEN, GREEN_LINE, size=16)
    add_line(slide, Inches(3.2), Inches(5.85), Inches(4.0), Inches(5.85))
    add_line(slide, Inches(6.25), Inches(5.85), Inches(7.05), Inches(5.85))
    add_footer(slide, "工业 AI 下一步不是继续堆模型，而是补系统能力")


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    for fn in [slide1, slide2, slide3, slide4, slide5, slide6, slide7, slide8, slide9, slide10, slide11, slide12]:
        fn(prs)

    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
